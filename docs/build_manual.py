from pathlib import Path
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import html


ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / "docs" / "Vertex_AI_Layout_to_JSON_Developer_User_Manual_FINAL.html"
DOCX = HTML.with_suffix(".docx")
PDF = HTML.with_suffix(".pdf")
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")


def set_cell_fill(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def add_rich_paragraph(doc, element, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(element.text_content())
    return paragraph


def add_table(doc, element) -> None:
    rows = element.xpath("./tr | ./thead/tr | ./tbody/tr")
    if not rows:
        return
    column_count = max(len(row.xpath("./th | ./td")) for row in rows)
    table = doc.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    for row_index, row_element in enumerate(rows):
        row = table.add_row()
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_properties.append(OxmlElement("w:tblHeader"))
        for column_index, cell_element in enumerate(row_element.xpath("./th | ./td")):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.add_run(cell_element.text_content())
            if row_index == 0 or cell_element.tag == "th":
                set_cell_fill(cell, "17365D")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_fill(cell, "F4F7FB")
    doc.add_paragraph()


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build_docx() -> None:
    root = html.fromstring(HTML.read_text(encoding="utf-8"))
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 26), ("Subtitle", 16), ("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Aptos Display" if "Heading" in name or name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False
        paragraph_properties = style.element.get_or_add_pPr()
        border = paragraph_properties.find(qn("w:pBdr"))
        if border is not None:
            paragraph_properties.remove(border)

    def emit(element, cover=False) -> None:
        if element.tag == "section":
            for child in element.iterchildren():
                emit(child, cover=True)
            doc.add_page_break()
        elif element.tag == "h1":
            paragraph = add_rich_paragraph(doc, element, "Title" if cover else "Heading 1")
            if "pagebreak" in element.get("class", "").split():
                paragraph.paragraph_format.page_break_before = True
            if cover:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif element.tag == "h2":
            paragraph = add_rich_paragraph(doc, element, "Subtitle" if cover else "Heading 2")
            if cover:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif element.tag == "h3":
            add_rich_paragraph(doc, element, "Heading 3")
        elif element.tag == "p":
            paragraph = add_rich_paragraph(doc, element)
            if cover:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif element.tag == "pre":
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(element.text_content())
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            paragraph.paragraph_format.space_after = Pt(8)
        elif element.tag in {"ul", "ol"}:
            start = int(element.get("start", "1"))
            for offset, item in enumerate(element.xpath("./li")):
                if element.tag == "ul":
                    add_rich_paragraph(doc, item, "List Bullet")
                else:
                    paragraph = doc.add_paragraph()
                    paragraph.paragraph_format.left_indent = Inches(0.25)
                    paragraph.paragraph_format.first_line_indent = Inches(-0.2)
                    paragraph.add_run(f"{start + offset}. {item.text_content()}")
        elif element.tag == "table":
            add_table(doc, element)
        elif element.tag == "div" and "shot" in element.get("class", "").split():
            images = element.xpath(".//img")
            captions = element.xpath(".//*[contains(concat(' ', normalize-space(@class), ' '), ' caption ')]")
            if images:
                image_path = HTML.parent / images[0].get("src")
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shape = paragraph.add_run().add_picture(str(image_path), width=Inches(6.8))
                shape._inline.docPr.set(
                    "descr",
                    captions[0].text_content() if captions else image_path.stem.replace("_", " "),
                )
            if captions:
                paragraph = doc.add_paragraph(captions[0].text_content(), style="Caption")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    body = root.find("body")
    for child in body.iterchildren():
        emit(child)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run("Vertex AI Layout to JSON Developer and User Manual    ")
    add_page_number(footer)
    doc.save(DOCX)


def build_pdf() -> None:
    if not EDGE.exists():
        raise FileNotFoundError(EDGE)
    subprocess.run(
        [
            str(EDGE),
            "--headless",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={PDF}",
            HTML.resolve().as_uri(),
        ],
        check=True,
    )


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(f"Built {DOCX}")
    print(f"Built {PDF}")
